# -*- coding: utf-8 -*-
"""将培养方案系列 markdown 转换为 Word(.docx)，精确保留表格与图片。
特性：标题/引用/水平线/GFM表格(含对齐与表头底纹)/有序无序列表(含嵌套)/行内加粗·代码/图片嵌入。
图片：md 内未内嵌图，按专业把对应的反向设计大图(PNG)作为"附图"嵌入，避免图片在 Word 版丢失。
"""
import os, re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Microsoft YaHei"
BODY = 10.5
HEAD_SIZES = {1: 16, 2: 14, 3: 12.5, 4: 11.5}
USABLE = Inches(6.27)  # A4 可用宽度

INLINE = re.compile(r'(\*\*.+?\*\*)|(`[^`]+?`)|(\*.+?\*)')


def set_run_font(r, name=FONT, size=BODY, bold=False, italic=False, code=False):
    r.font.name = name
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    rPr = r._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rPr.append(rf)
    rf.set(qn('w:eastAsia'), FONT if not code else 'Microsoft YaHei')
    rf.set(qn('w:ascii'), name)
    rf.set(qn('w:hAnsi'), name)


def add_inline(paragraph, text):
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            set_run_font(r)
        tok = m.group(0)
        if tok.startswith('**'):
            r = paragraph.add_run(tok[2:-2]); set_run_font(r, bold=True)
        elif tok.startswith('`'):
            r = paragraph.add_run(tok[1:-1]); set_run_font(r, name='Consolas', code=True)
        else:
            r = paragraph.add_run(tok[1:-1]); set_run_font(r, italic=True)
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:]); set_run_font(r)


def style_quote(p):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single'); left.set(qn('w:sz'), '18')
    left.set(qn('w:space'), '8'); left.set(qn('w:color'), '4472C4')
    pBdr.append(left); pPr.append(pBdr)
    p.paragraph_format.left_indent = Pt(14)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)


def add_hr(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '6')
    b.set(qn('w:space'), '1'); b.set(qn('w:color'), '808080')
    pBdr.append(b); pPr.append(pBdr)
    p.paragraph_format.space_after = Pt(2)


def add_heading(doc, text, level):
    h = doc.add_heading(level=level)
    # 清空默认 run 后由我们写入（支持行内格式）
    for r in list(h.runs):
        r._element.getparent().remove(r._element)
    add_inline(h, text)
    for r in h.runs:
        set_run_font(r, size=HEAD_SIZES[level], bold=True)
    return h


def split_row(line):
    s = line.strip()
    if s.startswith('|'):
        s = s[1:]
    if s.endswith('|'):
        s = s[:-1]
    return [p.strip() for p in s.split('|')]


def parse_align(sep):
    parts = split_row(sep)
    als = []
    for p in parts:
        p = p.strip()
        l = p.startswith(':'); r = p.endswith(':')
        if l and r:
            als.append('center')
        elif r:
            als.append('right')
        else:
            als.append('left')
    return als


def set_cell_text(cell, text, bold=False, align='left', size=9.5):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT,
                   'center': WD_ALIGN_PARAGRAPH.CENTER,
                   'right': WD_ALIGN_PARAGRAPH.RIGHT}[align]
    add_inline(p, text)
    for r in p.runs:
        set_run_font(r, size=size, bold=bold)


def build_table(doc, lines):
    header = split_row(lines[0])
    aligns = parse_align(lines[1])
    body = [split_row(l) for l in lines[2:]]
    ncol = len(header)
    t = doc.add_table(rows=1, cols=ncol)
    t.style = 'Table Grid'
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hdr = t.rows[0].cells
    for c in range(ncol):
        set_cell_text(hdr[c], header[c], bold=True, align=aligns[c], size=9.5)
        tcPr = hdr[c]._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), 'D9E2F3')
        tcPr.append(shd)
    for row in body:
        cells = t.add_row().cells
        for c in range(min(ncol, len(row))):
            set_cell_text(cells[c], row[c], align=aligns[c], size=9.5)
    for c in range(ncol):
        w = USABLE / ncol
        for row in t.rows:
            row.cells[c].width = w
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def append_figures(doc, figs):
    if not figs:
        return
    doc.add_page_break()
    add_heading(doc, '附图：反向设计大图', 2)
    for path, cap in figs:
        if not os.path.exists(path):
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(path, width=Inches(6.0))
        cap_p = doc.add_paragraph()
        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap_p.add_run(cap)
        set_run_font(r, size=9, italic=True)


def md_to_docx(md_path, docx_path, figs):
    with open(md_path, 'r', encoding='utf-8') as f:
        text = f.read()
    doc = Document()
    # A4 + 页边距
    sec = doc.sections[0]
    sec.page_height = Inches(11.69); sec.page_width = Inches(8.27)
    sec.left_margin = Inches(1.0); sec.right_margin = Inches(1.0)
    sec.top_margin = Inches(1.0); sec.bottom_margin = Inches(1.0)
    # 默认字体
    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(BODY)
    rPr = normal.element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rPr.append(rf)
    rf.set(qn('w:eastAsia'), FONT); rf.set(qn('w:ascii'), FONT); rf.set(qn('w:hAnsi'), FONT)

    lines = text.split('\n')
    in_code = False
    code_lines = []
    i = 0
    N = len(lines)
    while i < N:
        line = lines[i]
        stripped = line.strip()
        # 代码围栏
        if stripped.startswith('```'):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Pt(10)
                run = p.add_run('\n'.join(code_lines))
                set_run_font(run, name='Consolas', size=9, code=True)
                code_lines = []
                in_code = False
            else:
                in_code = True
                code_lines = []
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        # 表格
        if (stripped.startswith('|') and i + 1 < N
                and re.match(r'^\s*\|?[\s:|\-]+\|?\s*$', lines[i + 1].strip())
                and '-' in lines[i + 1]):
            tbl = [line]
            j = i + 1
            j += 1  # 跳过分隔行
            while j < N and lines[j].strip().startswith('|'):
                tbl.append(lines[j]); j += 1
            build_table(doc, tbl)
            i = j
            continue
        # 引用
        if stripped.startswith('>'):
            qbuf = []
            while i < N and lines[i].strip().startswith('>'):
                qbuf.append(lines[i].strip().lstrip('>').strip())
                i += 1
            p = doc.add_paragraph()
            add_inline(p, ' '.join(qbuf))
            style_quote(p)
            continue
        # 水平线
        if re.match(r'^(\-{3,}|\*{3,}|_{3,})$', stripped):
            add_hr(doc); i += 1; continue
        # 标题
        m = re.match(r'^(#{1,4})\s+(.*)$', line)
        if m:
            add_heading(doc, m.group(2), len(m.group(1))); i += 1; continue
        # 有序列表
        m = re.match(r'^(\s*)(\d+)\.\s+(.*)$', line)
        if m:
            lvl = len(m.group(1)) // 2
            st = 'List Number' if lvl == 0 else 'List Number ' + str(lvl + 1)
            p = doc.add_paragraph(style=st)
            add_inline(p, m.group(3)); i += 1; continue
        # 无序列表
        m = re.match(r'^(\s*)[-*]\s+(.*)$', line)
        if m:
            lvl = len(m.group(1)) // 2
            st = 'List Bullet' if lvl == 0 else 'List Bullet ' + str(lvl + 1)
            p = doc.add_paragraph(style=st)
            add_inline(p, m.group(2)); i += 1; continue
        # 空行
        if stripped == '':
            i += 1; continue
        # 普通段落
        p = doc.add_paragraph()
        add_inline(p, line)
        i += 1

    append_figures(doc, figs)
    doc.save(docx_path)


if __name__ == '__main__':
    base = r"F:/教学/2025年工作/智能科学与技术专业"
    jobs = [
        ("培养方案设计轮廓-三专业.md", "培养方案设计轮廓-三专业.docx",
         [(os.path.join(base, "reverse_cross.png"), "跨专业依赖总图（三专业系统 / 课程衔接关系）")]),
        ("人工智能专业-系统能力课程映射.md", "人工智能专业-系统能力课程映射.docx",
         [(os.path.join(base, "reverse_AI.png"), "人工智能专业反向设计大图（系统 → 子系统 → 功能 → 知识点 → 课程·大纲）")]),
        ("智能科学与技术专业-培养方案修订稿.md", "智能科学与技术专业-培养方案修订稿.docx",
         [(os.path.join(base, "reverse_ZK.png"), "智能科学与技术专业反向设计大图（10 大系统）"),
          (os.path.join(base, "reverse_cross.png"), "跨专业依赖总图")]),
    ]
    for md, docx, figs in jobs:
        mp = os.path.join(base, md)
        dp = os.path.join(base, docx)
        md_to_docx(mp, dp, figs)
        d = Document(dp)
        print(f"[OK] {docx} : 段落={len(d.paragraphs)} 表格={len(d.tables)}")
