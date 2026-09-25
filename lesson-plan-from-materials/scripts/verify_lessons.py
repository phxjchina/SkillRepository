#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
verify_lessons.py — 校验生成的教案 docx

用法:
    python verify_lessons.py <目录> [glob_pattern]
      默认 glob: *教案*.docx

对每个文件检查:
    - 能否被 python-docx 正常打开（捕获损坏/编码问题）
    - 表格数量是否 >= 2（课程信息表 + 教案主体表）
    - 打印表0/表1 中若干关键字段，便于人工抽检是否填对
"""
import sys
import glob
import os
from docx import Document


def main():
    if len(sys.argv) < 2:
        print("用法: python verify_lessons.py <目录> [glob]", file=sys.stderr)
        sys.exit(2)
    folder = sys.argv[1]
    pattern = sys.argv[2] if len(sys.argv) > 2 else "*教案*.docx"
    files = sorted(glob.glob(os.path.join(folder, pattern)))
    if not files:
        print(f"未找到匹配 {pattern} 的文件", file=sys.stderr)
        sys.exit(1)

    print(f"共 {len(files)} 个教案\n")
    ok = 0
    for p in files:
        name = os.path.basename(p)
        try:
            d = Document(p)
        except Exception as e:
            print(f"✗ {name}: 打开失败 -> {e}")
            continue
        ntbl = len(d.tables)
        status = "✓" if ntbl >= 2 else "⚠"
        if ntbl >= 2:
            ok += 1
        print(f"{status} {name}  (表格数 {ntbl})")
        if ntbl >= 2:
            t0 = d.tables[0]
            t1 = d.tables[1]

            def g0(r, c):
                try:
                    return t0.rows[r].cells[c].text.strip().replace("\n", " ")
                except Exception:
                    return ""

            def g1(r, c):
                try:
                    return t1.rows[r].cells[c].text.strip().replace("\n", " ")
                except Exception:
                    return ""

            # 课程信息表常见字段（按通用模板坐标，遇空则跳过）
            print(f"    课程名称: {g0(0,1)} | 课程编号: {g0(0,5)}")
            print(f"    授课教师: {g0(1,1)} | 班级: {g0(2,1)}")
            print(f"    教材: {g0(7,1)}")
            print(f"    授课题目: {g1(0,1)}")
            print(f"    授课进度: {g1(1,1)}")
            # 教学目标是否非空
            obj = g1(3, 1)
            print(f"    教学目标(前40字): {obj[:40]}{'…' if len(obj) > 40 else ''}")
        print()
    print(f"可用(>=2表): {ok}/{len(files)}")


if __name__ == "__main__":
    main()
