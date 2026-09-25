#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
analyze_template.py — 探查教案模板的表格结构

用法:
    python analyze_template.py <template.docx>

输出每个表格的 (行, 列) 坐标与单元格文本，便于：
    - 确认表0=课程信息表 / 表1=教案主体表 / 表2=教学内容表 的字段位置
    - 为 generate_lesson_docx.py 的 fills 坐标提供准确依据
    - 检查合并单元格（python-docx 合并区会重复返回同一 cell 引用，属正常）

注意：合并单元格在 row.cells 中同一区域会返回相同对象，填写其一即可。
"""
import sys
from docx import Document


def main():
    if len(sys.argv) < 2:
        print("用法: python analyze_template.py <template.docx>", file=sys.stderr)
        sys.exit(2)
    doc = Document(sys.argv[1])
    print(f"文档共有 {len(doc.tables)} 个表格\n")
    for ti, t in enumerate(doc.tables):
        print(f"==== 表 {ti}  ({len(t.rows)} 行 x {len(t.columns)} 列) ====")
        for ri, row in enumerate(t.rows):
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            print(f"R{ri}: " + " || ".join(cells))
        print()


if __name__ == "__main__":
    main()
