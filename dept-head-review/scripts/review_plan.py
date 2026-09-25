# -*- coding: utf-8 -*-
"""系主任审核 · 自动化初筛。
对一份或多份培养方案(.md/.docx)做文本级初筛：
  - 抽取课程名（《...》引号 + 像课程名的加粗词）
  - 多文件时检测跨专业重名课程（冗余风险）
  - 检测六维关键章节是否缺失（缺位预警）
  - 计数表格（课程清单/支撑矩阵多在表格里）
仅做"初筛"，终审仍按 dept-head-review SKILL 六维标准人工判断。
"""
import sys, re, os

try:
    from docx import Document
except Exception:
    Document = None

COURSE_RE = re.compile(r'《([^》]{2,40})》')
BOLD_COURSE_RE = re.compile(r'\*\*([^＊]{2,30}?)\*\*')
# 像课程名的加粗词：含这些关键字之一，且不太长
COURSE_HINT = ('课程', '基础', '系统', '设计', '算法', '原理', '技术', '应用', '导论', '数学',
               '视觉', '语言', '网络', '嵌入式', '智能', '学习', '模型', '工程', '概论', '实验')


def read_text(path):
    if path.lower().endswith('.docx'):
        if Document is None:
            return '[docx 解析不可用，请装 python-docx]'
        d = Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(' | '.join(c.text for c in row.cells))
        return '\n'.join(parts)
    with open(path, 'r', encoding='utf-8', errors='ignore') as f:
        return f.read()


def extract_courses(text):
    names = set()
    for m in COURSE_RE.finditer(text):
        names.add(m.group(1).strip())
    for m in BOLD_COURSE_RE.finditer(text):
        tok = m.group(1).strip()
        if any(h in tok for h in COURSE_HINT) and 2 < len(tok) < 24:
            names.add(tok)
    return names


def count_tables(path, text):
    if path.lower().endswith('.docx') and Document is not None:
        try:
            return len(Document(path).tables)
        except Exception:
            return 0
    return text.count('\n|---') + text.count('\n| ---') + len(re.findall(r'\n\|[-: |]+\|\n', text))


SECTION_KEYWORDS = {
    '差异化/错位': ('差异化', '错位', '一句话区分', '定位'),
    'OBE反向设计': ('反向设计', '能开发', '子系统', '知识点', '教学大纲'),
    '三层结构': ('基座课程', '核心课程', '特色课程', '三层'),
    '学期/学分': ('学期', '学分', 'S1', 'S3', 'S5', 'S7'),
    '撤销专业迁移': ('撤销', '迁移', '功能不撤销', '软工', '物联网'),
    '工程认证': ('毕业要求', '指标点', '认证', '支撑矩阵'),
}


def section_check(text):
    missing = []
    for dim, kws in SECTION_KEYWORDS.items():
        if not any(k in text for k in kws):
            missing.append(dim)
    return missing


def main():
    args = sys.argv[1:]
    if not args or '--help' in args:
        print("用法: python review_plan.py <方案1.md/.docx> [方案2 ...] [--compare]")
        return
    compare = '--compare' in args
    files = [a for a in args if not a.startswith('--')]

    data = {}
    for f in files:
        if not os.path.exists(f):
            print(f"[跳过] 文件不存在: {f}")
            continue
        text = read_text(f)
        data[f] = {
            'courses': extract_courses(text),
            'tables': count_tables(f, text),
            'missing': section_check(text),
            'chars': len(text),
        }

    print("# 培养方案审核 · 自动化初筛\n")
    for f, d in data.items():
        print(f"## {os.path.basename(f)}")
        print(f"- 字符数: {d['chars']}  表格数(估): {d['tables']}  抽取课程名: {len(d['courses'])}")
        print(f"- 缺失章节预警: {', '.join(d['missing']) if d['missing'] else '无（六维章节齐全）'}")
        cs = sorted(d['courses'])
        if cs:
            print(f"- 课程名抽样: {', '.join(cs[:25])}{' …' if len(cs) > 25 else ''}")
        print()

    if compare and len(data) >= 2:
        print("## 跨专业重名课程（冗余风险）")
        from collections import defaultdict
        occ = defaultdict(list)
        for f, d in data.items():
            for c in d['courses']:
                occ[c].append(os.path.basename(f))
        dups = {c: fs for c, fs in occ.items() if len(fs) > 1}
        if dups:
            for c, fs in sorted(dups.items()):
                print(f"- **{c}** 出现在: {', '.join(sorted(set(fs)))}")
        else:
            print("- 未发现跨文件重名课程（注意：仅基于抽取到的课程名，可能漏检近名课）")
        print("\n> 提示：近名不同字（如《深度学习》vs《深度学习A》）不会被判定为重名，需人工核对。")


if __name__ == '__main__':
    main()
